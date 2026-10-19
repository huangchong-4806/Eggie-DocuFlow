"""Local, immutable-preview Excel category splitting. Never edits the input."""
from __future__ import annotations

from copy import copy
from dataclasses import dataclass, replace
from datetime import date, datetime, time
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import re
import shutil
import warnings
import unicodedata

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.dimensions import RowDimension, ColumnDimension


@dataclass(frozen=True)
class SplitOptions:
    source: Path
    sheet: str
    area: str | None = None
    header_rows: int = 1
    key_columns: tuple[int, ...] = ()
    trim_keys: bool = False
    visible_only: bool = False
    output_mode: str = 'files'
    output_dir: Path = Path('.')


@dataclass(frozen=True)
class SourceRow:
    number: int
    values: tuple[object, ...]
    hidden: bool
    blank: bool


@dataclass(frozen=True)
class SourceSnapshot:
    source_sha256: str
    data_rows: tuple[SourceRow, ...]
    has_filter: bool
    has_hidden: bool
    issues: tuple[str, ...]
    options: SplitOptions
    bounds: tuple[int, int, int, int]
    headers: tuple[SourceRow, ...]


@dataclass(frozen=True)
class Group:
    id: int
    label: str
    source_rows: tuple[int, ...]


@dataclass(frozen=True)
class Counts:
    total: int
    hidden: int
    blank: int
    unselected: int
    exported: int


@dataclass(frozen=True)
class SplitPreview:
    options: SplitOptions
    source_sha256: str
    groups: tuple[Group, ...]
    selected_ids: tuple[int, ...]
    counts: Counts
    names: tuple[str, ...]
    result_folder: Path
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True)
class ExportResult:
    output_files: tuple[Path, ...]
    exported_rows: int
    row_map: tuple[tuple[int, str, str, int], ...]
    output_folder: Path


def _read_bytes(source):
    source = Path(source)
    if source.suffix.lower() != '.xlsx':
        raise ValueError('分类拆分目前只支持 .xlsx 文件。')
    if not source.is_file():
        raise ValueError('源文件不存在，请重新选择。')
    return source.read_bytes()


def _open(data, **kwargs):
    # Unsupported Excel extension warnings must never silently lose content.
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always', UserWarning)
        wb = load_workbook(BytesIO(data), keep_links=False, **kwargs)
    if caught:
        wb.close()
        raise ValueError('文件包含暂不能可靠保留的 Excel 扩展内容，请先另存普通明细表。')
    return wb


def source_info(source, sheet=None, area=None, header_rows=1):
    """Return sheet names and actual source-column labels without changing data."""
    data = _read_bytes(source)
    wb = _open(data, read_only=True, data_only=True)
    try:
        names = tuple(wb.sheetnames)
        ws = wb[sheet or names[0]]
        bounds = _bounds(ws, area, header_rows)
        left, top, right, bottom = bounds
        labels = []
        values = next(ws.iter_rows(min_row=top + max(header_rows-1, 0),
                                   max_row=top + max(header_rows-1, 0),
                                   min_col=left, max_col=right, values_only=True))
        for col, value in zip(range(left, right+1), values):
            label = f'{get_column_letter(col)}列'
            if header_rows and value is not None:
                label += f' — {value}'
            labels.append((col, label))
        return names, tuple(labels), sha256(data).hexdigest()
    finally:
        wb.close()


def _bounds(ws, area, header_rows):
    if area:
        if not re.fullmatch(r'\$?[A-Za-z]{1,3}\$?[1-9]\d*:\$?[A-Za-z]{1,3}\$?[1-9]\d*', area.strip()):
            raise ValueError('区域请填写完整矩形，例如 A3:H500。')
        bounds = range_boundaries(area.strip())
    else:
        bounds = (1, 1, ws.max_column or 1, ws.max_row or 1)
    left, top, right, bottom = bounds
    if not (1 <= left <= right <= 16384 and 1 <= top <= bottom <= 1048576):
        raise ValueError('区域超出 Excel 范围或起止顺序不正确。')
    if header_rows < 0 or header_rows >= bottom-top+1:
        raise ValueError('表头之后没有可拆分数据，请检查表头行数和区域。')
    # Bound memory before normal worksheet iteration creates empty cells.
    if (right-left+1)*(bottom-top+1) > 2_000_000:
        raise ValueError('所选区域超过 200 万单元格，请缩小处理区域。')
    return bounds


def _overlap(a, b):
    return a[0] <= b[2] and b[0] <= a[2] and a[1] <= b[3] and b[1] <= a[3]


def _filter_active(af, bounds):
    if af is None or not af.ref or not af.filterColumn:
        return False
    return _overlap(range_boundaries(af.ref), bounds) and any(
        any(getattr(c, kind, None) is not None for kind in
            ('filters','customFilters','dynamicFilter','top10','colorFilter','iconFilter'))
        for c in af.filterColumn
    )


def inspect_source(options):
    data = _read_bytes(options.source)
    wb = _open(data, data_only=False)
    cached = None
    try:
        if options.sheet not in wb.sheetnames:
            raise ValueError('工作表不存在，请重新读取文件。')
        ws = wb[options.sheet]
        bounds = _bounds(ws, options.area, options.header_rows)
        left, top, right, bottom = bounds
        if any(c < left or c > right for c in options.key_columns):
            raise ValueError('分类列必须位于所选区域内。')
        if ws._charts or ws._images or ws._pivots or len(ws.conditional_formatting) or ws.data_validations.count:
            raise ValueError(f'{ws.title} 包含图表、图片、透视表、条件格式或数据验证，当前不能完整保留；请先另存普通明细表。')
        cached = _open(data, data_only=True)
        value_ws = cached[options.sheet]
        issues = []
        merged_sources = {}
        header_end = top + options.header_rows - 1
        for merged in ws.merged_cells.ranges:
            box = merged.bounds
            if not _overlap(bounds, box):
                continue
            a, b, c, d = box
            if a < left or c > right or b < top or d > bottom or b <= header_end < d:
                issues.append(f'{ws.title}!{merged} 跨越处理区域或表头边界，请先调整。')
                continue
            if b > header_end and d > b:
                if c > a:
                    issues.append(f'{ws.title}!{merged} 同时跨多行多列，请先取消合并。')
                    continue
                anchor = ws.cell(b, a)
                if a not in options.key_columns and (anchor.data_type == 'f' or isinstance(anchor.value, (int,float,datetime,date,time))):
                    issues.append(f'{ws.title}!{merged} 是非分类列合并数字或公式，需先明确分配，避免重复累计。')
                    continue
                for row in range(b, d+1):
                    merged_sources[(row,a)] = (b,a)
        has_filter = _filter_active(ws.auto_filter, bounds) or any(
            _filter_active(t.autoFilter, bounds) for t in ws.tables.values())
        headers, rows = [], []
        for row in range(top, bottom+1):
            hidden = bool(ws.row_dimensions[row].hidden)
            included = row <= header_end or not options.visible_only or not hidden
            values = []
            for col in range(left,right+1):
                rr, cc = merged_sources.get((row,col),(row,col))
                cell = ws.cell(rr,cc)
                value = cell.value
                if cell.data_type == 'f':
                    result = value_ws.cell(rr,cc)
                    if included and (result.value is None or result.data_type == 'e'):
                        issues.append(f'{ws.title}!{cell.coordinate} 公式缺少已保存结果或结果错误，请用 Excel 计算并保存后重新读取。')
                    value = result.value
                values.append(value)
            item = SourceRow(row,tuple(values),hidden,all(v is None or v == '' for v in values))
            (headers if row <= header_end else rows).append(item)
        return SourceSnapshot(sha256(data).hexdigest(),tuple(rows),bool(has_filter),
                              any(r.hidden for r in rows),tuple(dict.fromkeys(issues)),options,bounds,tuple(headers))
    finally:
        wb.close()
        if cached is not None:
            cached.close()


def _key(value, trim=False):
    if isinstance(value,str) and trim:
        value=value.strip()
    if value is None or value == '':
        return ('blank', '')
    if isinstance(value,datetime):
        return ('datetime',value.isoformat())
    if isinstance(value,(date,time)):
        return (type(value).__name__,value.isoformat())
    if isinstance(value,bool):
        return ('bool',value)
    if isinstance(value,(float,int)):
        return ('number',value)
    return (type(value).__name__,value)


def _label(key):
    return '_'.join('未分类' if kind == 'blank' else str(value) for kind,value in key)


def _safe_name(text, limit):
    text = re.sub(r'[\x00-\x1f<>:"/\\|?*\[\]]','_',str(text)).strip(" .'") or '未分类'
    if re.fullmatch(r'(?i:CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(?:\..*)?',text):
        text = '_' + text
    # UTF-8 byte ceiling also protects long Chinese file names on macOS.
    while len(text.encode('utf-8')) > 180:
        text=text[:-1]
    return _truncate_name(text, limit).rstrip(" .'") or '未分类'


def _truncate_name(text, limit):
    while len(text.encode('utf-16-le')) // 2 > limit:
        text = text[:-1]
    return text


def _copy_style(origin, target):
    for name in ('font', 'fill', 'border', 'alignment', 'protection'):
        setattr(target, name, copy(getattr(origin, name)))
    target.number_format = origin.number_format


def _names(groups, options):
    used=set(); names=[]
    stem=_safe_name(Path(options.source).stem,50)
    for g in groups:
        raw=g.label if options.output_mode=='sheets' else f'{stem}_{g.label}'
        limit=31 if options.output_mode=='sheets' else 100
        base=_safe_name(raw,limit)
        if options.output_mode=='sheets' and base.casefold()=='history':
            base='_'+base
        candidate=base; number=2
        while unicodedata.normalize('NFD', candidate).casefold() in used:
            suffix=f'_{number}'
            candidate=_truncate_name(base,limit-len(suffix))+suffix
            number+=1
        used.add(unicodedata.normalize('NFD', candidate).casefold()); names.append(candidate if options.output_mode=='sheets' else candidate+'.xlsx')
    return tuple(names)


def build_preview(snapshot, options):
    if snapshot.options != options:
        raise ValueError('设置已改变，请重新读取预览。')
    if snapshot.issues:
        raise ValueError('\n'.join(snapshot.issues[:30]) + (f'\n另有 {len(snapshot.issues)-30} 项。' if len(snapshot.issues)>30 else ''))
    if not options.key_columns or len(set(options.key_columns))!=len(options.key_columns):
        raise ValueError('请至少选择一个分类列，不能重复选择。')
    if options.output_mode not in ('files','sheets'):
        raise ValueError('输出方式不正确。')
    parent=Path(options.output_dir).resolve()
    if not parent.is_dir():
        raise ValueError('请选择已存在的保存文件夹。')
    buckets={}; hidden=blank=0; whitespace=False
    for row in snapshot.data_rows:
        if options.visible_only and row.hidden:
            hidden+=1; continue
        if row.blank:
            blank+=1; continue
        vals=[row.values[c-snapshot.bounds[0]] for c in options.key_columns]
        whitespace |= any(isinstance(v,str) and v!=v.strip() for v in vals)
        key=tuple(_key(v,options.trim_keys) for v in vals)
        buckets.setdefault(key,[]).append(row.number)
    groups=tuple(Group(i,_label(k),tuple(rows)) for i,(k,rows) in enumerate(buckets.items()))
    if not groups:
        raise ValueError('当前范围没有可拆分的非空数据。')
    base=_safe_name(Path(options.source).stem,60)+'_分类拆分结果'
    folder=parent/base; n=1
    while folder.exists():
        folder=parent/f'{base}_{n}'; n+=1
    counts=Counts(len(snapshot.data_rows),hidden,blank,0,sum(len(g.source_rows) for g in groups))
    return SplitPreview(options,snapshot.source_sha256,groups,tuple(g.id for g in groups),counts,
                        _names(groups,options),folder,
                        ('发现分类名称含首尾空格，可返回设置选择忽略首尾空格。',) if whitespace else ())


def select_groups(preview, ids):
    wanted=set(ids)
    if not wanted.issubset({g.id for g in preview.groups}):
        raise ValueError('所选类别已失效，请重新预览。')
    selected=tuple(g.id for g in preview.groups if g.id in wanted)
    total=sum(len(g.source_rows) for g in preview.groups)
    exported=sum(len(g.source_rows) for g in preview.groups if g.id in wanted)
    return replace(preview,selected_ids=selected,counts=replace(preview.counts,exported=exported,unselected=total-exported))


def plan_names(preview):
    return tuple(preview.names[g.id] for g in preview.groups if g.id in preview.selected_ids)


def export_split(preview, progress_callback=None):
    if not preview.selected_ids:
        raise ValueError('请至少选择一个类别。')
    options=preview.options
    data=_read_bytes(options.source)
    if sha256(data).hexdigest()!=preview.source_sha256:
        raise ValueError('源文件已变化，请重新读取、预览并确认。')
    snapshot=inspect_source(options)
    if snapshot.source_sha256!=preview.source_sha256:
        raise ValueError('源文件已变化，请重新预览。')
    validated=build_preview(snapshot,options)
    if validated.groups!=preview.groups or validated.names!=preview.names:
        raise ValueError('分类或名称已变化，请重新预览。')
    if snapshot.issues:
        raise ValueError('\n'.join(snapshot.issues))
    source=_open(data,data_only=False)
    folder=preview.result_folder
    created=False; files=[]; row_map=[]; output=None; expected_rows={}
    try:
        try:
            folder.mkdir()  # exclusive ownership; never delete a pre-existing directory
            created=True
        except FileExistsError:
            raise ValueError('结果位置已被占用，请重新预览确认新的保存位置。') from None
        ws=source[options.sheet]
        left,top,right,bottom=snapshot.bounds
        rows={r.number:r for r in snapshot.data_rows}
        selected=[g for g in preview.groups if g.id in preview.selected_ids]
        if options.output_mode=='sheets':
            output=Workbook(); output.loaded_theme=source.loaded_theme; output.remove(output.active)
        for index,g in enumerate(selected):
            if options.output_mode=='files':
                output=Workbook(); output.loaded_theme=source.loaded_theme; target=output.active; target.title='分类结果'
                path=folder/preview.names[g.id]
            else:
                target=output.create_sheet(preview.names[g.id])
                path=folder/(_safe_name(Path(options.source).stem,70)+'_分类结果.xlsx')
            mapping={}
            expected_rows[(str(path),target.title)] = tuple(snapshot.headers) + tuple(rows[n] for n in g.source_rows)
            for out_row,item in enumerate((*snapshot.headers,*(rows[n] for n in g.source_rows)),1):
                mapping[item.number]=out_row
                origin_dimension=ws.row_dimensions[item.number]
                dimension=RowDimension(target,index=out_row,height=origin_dimension.height)
                if origin_dimension.has_style:
                    _copy_style(origin_dimension,dimension)
                target.row_dimensions[out_row]=dimension
                for source_col,value in zip(range(left,right+1),item.values):
                    origin=ws.cell(item.number,source_col)
                    cell=target.cell(out_row,source_col-left+1,value)
                    if isinstance(value,str):
                        cell.data_type='s'
                    if origin.has_style:
                        _copy_style(origin,cell)
                if item.number >= top + options.header_rows:
                    row_map.append((item.number,str(path),target.title,out_row))
            for col in range(left,right+1):
                letter=get_column_letter(col-left+1)
                origin_dimension=next((dim for dim in ws.column_dimensions.values()
                                       if (dim.min or 0)<=col<=(dim.max or 0)),None)
                if origin_dimension is None:
                    origin_dimension=ws.column_dimensions[get_column_letter(col)]
                dimension=ColumnDimension(target,index=letter,width=origin_dimension.width,
                                          hidden=origin_dimension.hidden,min=col-left+1,max=col-left+1)
                if origin_dimension.has_style:
                    _copy_style(origin_dimension,dimension)
                target.column_dimensions[letter]=dimension
            for merge in ws.merged_cells.ranges:
                a,b,c,d=merge.bounds
                if a<left or c>right or b<top or d>bottom:
                    continue
                is_header=d<top+options.header_rows
                if (is_header or b==d) and all(r in mapping for r in range(b,d+1)):
                    target.merge_cells(start_row=mapping[b],end_row=mapping[d],start_column=a-left+1,end_column=c-left+1)
            if options.header_rows:
                target.freeze_panes=f'A{options.header_rows+1}'
            if options.output_mode=='files':
                output.save(path); output.close();output=None;files.append(path)
            if progress_callback:
                progress_callback(index+1,len(selected),f'正在生成：{preview.names[g.id]}')
        if options.output_mode=='sheets':
            output.save(path);output.close();output=None;files.append(path)
        if sha256(_read_bytes(options.source)).hexdigest()!=preview.source_sha256:
            raise ValueError('处理期间源文件已变化，本次结果已撤回，请重新预览。')
        # Independently read actual outputs before announcing success.
        actual=0
        for path in files:
            check=load_workbook(path,data_only=False)
            try:
                for sheet in check:
                    actual+=sheet.max_row-options.header_rows
                    expected=expected_rows[(str(path),sheet.title)]
                    for row,item in zip(sheet.iter_rows(min_col=1,max_col=right-left+1),expected):
                        if any(c.data_type=='f' for c in row):
                            raise ValueError('结果中出现未转换的公式，已停止交付。')
                        actual_values=tuple(None if c.value=='' else c.value for c in row)
                        expected_values=tuple(None if v=='' else v for v in item.values)
                        if actual_values!=expected_values:
                            raise ValueError(f'{sheet.title} 第 {row[0].row} 行内容核对不一致，已停止交付。')
            finally:
                check.close()
        if actual!=preview.counts.exported or len(row_map)!=actual:
            raise ValueError('导出行数核对不一致，已停止交付。')
        return ExportResult(tuple(files),actual,tuple(row_map),folder)
    except Exception:
        if created:
            shutil.rmtree(folder,ignore_errors=True)
        raise
    finally:
        source.close()
        if output is not None:
            output.close()
