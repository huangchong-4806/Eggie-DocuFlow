import hashlib
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.worksheet.table import Table

from excel_category_split import (
    SplitOptions, inspect_source, build_preview, select_groups, export_split,
)


class CategorySplitTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.source = self.root / '订单.xlsx'

    def make(self, rows, configure=None):
        wb = Workbook()
        ws = wb.active
        ws.title = '明细'
        for row in rows:
            ws.append(row)
        if configure:
            configure(ws)
        wb.save(self.source)
        wb.close()
        return SplitOptions(self.source, '明细', key_columns=(1,), output_dir=self.root)

    def preview(self, opts):
        return build_preview(inspect_source(opts), opts)

    def test_noncontiguous_groups_and_duplicate_rows_are_preserved(self):
        opts = self.make([['供应商', '金额'], ['A', 10], ['B', 20], ['A', 10], [None, 5]])
        p = self.preview(opts)
        self.assertEqual([g.source_rows for g in p.groups], [(2, 4), (3,), (5,)])
        self.assertEqual(p.counts.exported, 4)
        p = select_groups(p, (0,))
        self.assertEqual((p.counts.exported, p.counts.unselected), (2, 2))
        before = self.source.read_bytes()
        result = export_split(p)
        self.assertEqual(self.source.read_bytes(), before)
        wb = load_workbook(result.output_files[0])
        self.addCleanup(wb.close)
        self.assertEqual(list(wb.active.values), [('供应商', '金额'), ('A', 10), ('A', 10)])
        self.assertEqual([r[0] for r in result.row_map], [2, 4])

    def test_hidden_and_blank_counts_are_disjoint(self):
        def setup(ws):
            ws.row_dimensions[3].hidden = True
            ws.auto_filter.ref = 'A1:B5'
            ws.auto_filter.add_filter_column(0, ['A'])
        opts = self.make([['类别','金额'], ['A',1], [None,None], [None,None], ['B',2]], setup)
        snapshot = inspect_source(opts)
        self.assertTrue(snapshot.has_filter)
        self.assertTrue(snapshot.has_hidden)
        opts = replace(opts, visible_only=True)
        p = self.preview(opts)
        self.assertEqual((p.counts.total,p.counts.hidden,p.counts.blank,p.counts.exported), (4,1,1,2))

    def test_filter_arrows_alone_do_not_trigger(self):
        opts = self.make([['类别'],['A']], lambda ws: setattr(ws.auto_filter, 'ref', 'A1:A2'))
        self.assertFalse(inspect_source(opts).has_filter)

    def test_table_filter_is_detected(self):
        def setup(ws):
            af = AutoFilter(ref='A1:B3')
            af.add_filter_column(0, ['A'])
            ws.add_table(Table(displayName='Orders', ref='A1:B3', autoFilter=af))
        opts = self.make([['类别','金额'], ['A',1], ['B',2]],setup)
        self.assertTrue(inspect_source(opts).has_filter)

    def test_strict_keys_trim_and_blank_label_do_not_collide(self):
        opts = self.make([['类别'], ['A'], ['A '], ['001'], [1], ['1'], [None], ['未分类']])
        # Blank-only row is skipped, unlike a row with other data.
        self.assertEqual(len(self.preview(opts).groups),6)
        opts = replace(opts, trim_keys=True)
        p = self.preview(opts)
        self.assertEqual(len(p.groups),5)
        self.assertTrue(p.warnings)
        result = export_split(select_groups(p, (0,)))
        wb = load_workbook(result.output_files[0]); self.addCleanup(wb.close)
        self.assertEqual(wb.active['A3'].value, 'A ')

    def test_combination_keys_are_not_joined_strings(self):
        opts = self.make([['供应商','地区','值'], ['A_B','C',1], ['A','B_C',2], [None,'C',3], ['未分类','C',4]])
        p = self.preview(replace(opts,key_columns=(1,2)))
        self.assertEqual(len(p.groups),4)
        self.assertEqual(len({n.casefold() for n in p.names}),4)

    def test_region_headers_and_horizontal_merges_styles(self):
        def setup(ws):
            ws.merge_cells('B2:C2'); ws['B2']='标题'
            ws['B2'].font=Font(bold=True)
            ws.column_dimensions['B'].width=22
            ws.row_dimensions[4].height=32
            ws['C4'].number_format='0.00'
        opts = self.make([['外部'],[None,'标题',None],[None,'供应商','金额'],[None,'A',12],[None,'B',23]],setup)
        opts=replace(opts,area='B2:C5',header_rows=2,key_columns=(2,),output_mode='sheets')
        result=export_split(self.preview(opts))
        wb=load_workbook(result.output_files[0]);self.addCleanup(wb.close)
        self.assertEqual(wb.sheetnames,['A','B'])
        ws=wb['A']; self.assertIn('A1:B1',ws.merged_cells)
        self.assertTrue(ws['A1'].font.bold)
        self.assertEqual(ws.column_dimensions['A'].width,22)
        self.assertEqual(ws.row_dimensions[3].height,32)
        self.assertEqual(ws['B3'].number_format,'0.00')

    def test_vertical_category_merge_expands_and_numeric_other_merge_blocks(self):
        opts=self.make([['类别','金额'],['A',1],[None,2]],lambda ws:ws.merge_cells('A2:A3'))
        p=self.preview(opts);self.assertEqual(p.groups[0].source_rows,(2,3))
        result=export_split(p)
        wb=load_workbook(result.output_files[0]);self.addCleanup(wb.close)
        self.assertEqual(wb.active['A3'].value,'A')
        opts=self.make([['类别','金额'],['A',1000],['B',None]],lambda ws:ws.merge_cells('B2:B3'))
        with self.assertRaisesRegex(ValueError,'B2:B3'):
            self.preview(opts)

    def test_boundary_and_two_dimensional_merges_block(self):
        opts=self.make([['类别','金额'],['A',1],[None,None]],lambda ws:ws.merge_cells('A2:B3'))
        for area in (None,'B1:B3'):
            with self.subTest(area=area), self.assertRaises(ValueError):
                self.preview(replace(opts,area=area,key_columns=(2,) if area else (1,)))

    def test_missing_formula_cache_blocks_including_blank_looking_rows(self):
        opts=self.make([['类别','金额'],['A','=1+2']])
        with self.assertRaisesRegex(ValueError,'B2'):
            self.preview(opts)

    def test_literal_equals_text_is_not_formula(self):
        def setup(ws):
            ws['B2']='=文字';ws['B2'].data_type='s'
        opts=self.make([['类别','备注'],['A',None]],setup)
        result=export_split(self.preview(opts))
        wb=load_workbook(result.output_files[0]);self.addCleanup(wb.close)
        self.assertEqual(wb.active['B2'].data_type,'s')
        self.assertEqual(wb.active['B2'].value,'=文字')

    def test_safe_names_cannot_overwrite_or_merge_groups(self):
        opts=self.make([['类别'],['A/B'],['A:B'],['A_B_2'],['a_b'],['CON'],['x'*160]])
        p=self.preview(opts)
        self.assertEqual(len(set(n.casefold() for n in p.names)),6)
        r=export_split(p)
        self.assertEqual(len(r.output_files),6)
        for f in r.output_files:self.assertTrue(f.exists())
        with self.assertRaises(ValueError):export_split(p)
        self.assertTrue(all(f.exists() for f in r.output_files))

    def test_unicode_equivalent_names_remain_separate_files(self):
        import unicodedata
        opts=self.make([['类别','金额'],['é',10],['e\u0301',20]])
        p=self.preview(opts)
        self.assertEqual(len({unicodedata.normalize('NFD',n).casefold() for n in p.names}),2)
        r=export_split(p)
        self.assertEqual(len(r.output_files),2)
        actual=[]
        for path in r.output_files:
            wb=load_workbook(path)
            actual.append(tuple(wb.active.values)[1])
            wb.close()
        self.assertEqual(actual,[('é',10),('e\u0301',20)])

    def test_source_change_and_empty_selection_block_without_output(self):
        opts=self.make([['类别'],['A']]);p=self.preview(opts)
        with self.assertRaises(ValueError):export_split(select_groups(p,()))
        self.make([['类别'],['B']])
        with self.assertRaisesRegex(ValueError,'变化|改变'):export_split(p)
        self.assertEqual(list(self.root.iterdir()),[self.source])

    def test_save_failure_removes_only_owned_result(self):
        opts=self.make([['类别'],['A'],['B']]);p=self.preview(opts)
        with patch('openpyxl.workbook.workbook.Workbook.save',side_effect=OSError('disk full')):
            with self.assertRaises(OSError):export_split(p)
        self.assertEqual(list(self.root.iterdir()),[self.source])

    def test_options_change_cannot_reuse_snapshot(self):
        opts=self.make([['类别'],['A']]);snap=inspect_source(opts)
        with self.assertRaises(ValueError):build_preview(snap,replace(opts,header_rows=0))

    def test_invalid_area_and_keys_rejected(self):
        opts=self.make([['类别'],['A']])
        for changed in (replace(opts,area='A:A'),replace(opts,area='A1:B9999999'),replace(opts,key_columns=(2,)),replace(opts,header_rows=2)):
            with self.subTest(changed=changed),self.assertRaises(ValueError):self.preview(changed)

    def test_excel_calculated_cache_exports_values_and_formula_keys(self):
        source=Path(__file__).parent/'fixtures/category_formula_cache.xlsx'
        opts=SplitOptions(source,'公式明细',key_columns=(1,),output_dir=self.root,output_mode='sheets')
        result=export_split(self.preview(opts))
        wb=load_workbook(result.output_files[0],data_only=False);self.addCleanup(wb.close)
        self.assertEqual(wb['A供应商']['D2'].value,25)
        self.assertEqual(wb['A供应商']['D3'].value,0)
        self.assertIs(wb['A供应商']['E3'].value,False)
        self.assertEqual(wb['B供应商']['E2'].value,'=原始文字')
        self.assertEqual(wb['B供应商']['E2'].data_type,'s')
        for ws in wb:
            self.assertFalse(any(c.data_type=='f' for row in ws for c in row))
        self.assertEqual([g.label for g in self.preview(replace(opts,sheet='分类公式')).groups],['A','B'])
        with self.assertRaisesRegex(ValueError,'B2'):
            self.preview(replace(opts,sheet='错误结果'))

    def test_row_and_column_styles_are_registered_in_new_workbook(self):
        def setup(ws):
            ws.row_dimensions[2].font=Font(name='Arial',bold=True)
            ws.column_dimensions['B'].font=Font(name='Arial',italic=True)
        opts=self.make([['类别','值'],['A',1]],setup)
        r=export_split(self.preview(opts))
        wb=load_workbook(r.output_files[0]);self.addCleanup(wb.close)
        self.assertTrue(wb.active.row_dimensions[2].font.bold)
        self.assertTrue(wb.active.column_dimensions['B'].font.italic)

    def test_emoji_sheet_names_fit_excel_utf16_limit(self):
        opts=self.make([['类别'],['😀'*30]])
        p=self.preview(replace(opts,output_mode='sheets'))
        self.assertLessEqual(len(p.names[0].encode('utf-16-le'))//2,31)

    def test_zero_headers_region_rebases_columns_without_dropping_first_record(self):
        opts=self.make([['outside','A',1],['outside','B',2],['outside','A',3]])
        p=self.preview(replace(opts,area='B1:C3',header_rows=0,key_columns=(2,),output_mode='sheets'))
        r=export_split(p);w=load_workbook(r.output_files[0]);self.addCleanup(w.close)
        self.assertEqual(list(w['A'].values),[('A',1),('A',3)])
        self.assertEqual(r.exported_rows,3)

    def test_grouped_column_widths_preserved(self):
        def setup(ws):
            ws.column_dimensions.group('B','D',hidden=False)
            ws.column_dimensions['B'].width=27
        opts=self.make([['类别','甲','乙','丙'],['A',1,2,3]],setup)
        r=export_split(self.preview(opts));w=load_workbook(r.output_files[0]);self.addCleanup(w.close)
        self.assertEqual(w.active.column_dimensions['D'].width,27)

    def test_workbook_theme_is_preserved(self):
        opts=self.make([['类别'],['A']])
        w=load_workbook(self.source)
        w.loaded_theme=w.loaded_theme.replace(b'4F81BD',b'AA0000')
        theme=w.loaded_theme;w.save(self.source);w.close()
        r=export_split(self.preview(opts));w=load_workbook(r.output_files[0]);self.addCleanup(w.close)
        self.assertEqual(w.loaded_theme,theme)

    def test_reserved_history_sheet_name_is_made_safe(self):
        opts=self.make([['类别'],['History'],['history']])
        p=self.preview(replace(opts,output_mode='sheets'))
        self.assertNotIn('history',[n.casefold() for n in p.names])

    def test_cancelled_visibility_excludes_hidden_formula_errors(self):
        def setup(ws):ws.row_dimensions[3].hidden=True
        opts=self.make([['类别','金额'],['A',1],['B','=1/0']],setup)
        p=self.preview(replace(opts,visible_only=True))
        self.assertEqual(p.counts.exported,1)
        self.assertEqual(p.counts.hidden,1)

    def test_mid_export_source_change_removes_new_outputs(self):
        opts=self.make([['类别'],['A'],['B']]);p=self.preview(opts)
        def progress(*args):self.source.write_bytes(self.source.read_bytes()+b'changed')
        with self.assertRaisesRegex(ValueError,'变化'):export_split(p,progress)
        self.assertFalse(p.result_folder.exists())

    def test_content_corruption_during_save_is_detected(self):
        opts=self.make([['类别','金额'],['A',42]]);p=self.preview(opts)
        original=Workbook.save
        def corrupt(wb,path):
            wb.active['B2']=999
            original(wb,path)
        with patch.object(Workbook,'save',corrupt):
            with self.assertRaisesRegex(ValueError,'核对'):export_split(p)
        self.assertFalse(p.result_folder.exists())

    def test_complex_objects_block_instead_of_disappearing(self):
        from openpyxl.worksheet.datavalidation import DataValidation
        def setup(ws):
            dv=DataValidation(type='list',formula1='"A,B"')
            dv.add('A2');ws.add_data_validation(dv)
        opts=self.make([['类别'],['A']],setup)
        with self.assertRaisesRegex(ValueError,'验证|保留'):self.preview(opts)

    def test_randomized_multicolumn_splits_preserve_every_selected_source_row(self):
        import random
        for seed in range(12):
            rng=random.Random(seed)
            rows=[['供应商','区域','编号','数量']]
            for n in range(40):
                rows.append([rng.choice(['A','B','A ',None]),rng.choice(['东','西',None]),f'{n:04}',rng.choice([0,1,False,2.5])])
            hidden=set(rng.sample(range(2,42),8))
            def setup(ws):
                for n in hidden:ws.row_dimensions[n].hidden=True
            opts=self.make(rows,setup)
            opts=replace(opts,key_columns=(1,2),trim_keys=bool(seed%2),visible_only=bool(seed%3),output_mode='sheets' if seed%2 else 'files')
            p=self.preview(opts)
            selected=tuple(g.id for g in p.groups if g.id%2==0)
            p=select_groups(p,selected);r=export_split(p)
            mapped={n:(path,sheet,outrow) for n,path,sheet,outrow in r.row_map}
            expected={n for g in p.groups if g.id in selected for n in g.source_rows}
            self.assertEqual(set(mapped),expected)
            self.assertEqual(len(r.row_map),len(expected))
            actual={}
            for f in r.output_files:
                w=load_workbook(f)
                try:
                    for ws in w:
                        for row in ws.iter_rows(min_row=2,values_only=True):actual[row[2]]=row
                finally:w.close()
            for n in expected:
                self.assertEqual(actual[rows[n-1][2]],tuple(rows[n-1]))
            self.assertEqual(p.counts.total,p.counts.hidden+p.counts.blank+p.counts.unselected+p.counts.exported)

    def test_all_blank_or_all_hidden_selection_cannot_export(self):
        opts=self.make([['类别','值'],[None,None],[None,None]])
        with self.assertRaises(ValueError):self.preview(opts)
        opts=self.make([['类别'],['A']],lambda ws:setattr(ws.row_dimensions[2],'hidden',True))
        with self.assertRaises(ValueError):self.preview(replace(opts,visible_only=True))

    def test_duplicate_selection_ids_do_not_duplicate_exports(self):
        opts=self.make([['类别'],['A'],['B']]);p=select_groups(self.preview(opts),(0,0))
        self.assertEqual(p.counts.exported,1)
        r=export_split(p);self.assertEqual(r.exported_rows,1)

    def test_filter_outside_selected_range_does_not_affect_prompt(self):
        def setup(ws):
            ws.auto_filter.ref='C1:D3';ws.auto_filter.add_filter_column(0,['A'])
        opts=self.make([['类别','值','其他','值'],['A',1,'A',1],['B',2,'B',2]],setup)
        self.assertFalse(inspect_source(replace(opts,area='A1:B3')).has_filter)

    def test_text_vertical_merge_is_repeated_without_changing_source(self):
        opts=self.make([['类别','订单号'],['A','001'],['B',None]],lambda ws:ws.merge_cells('B2:B3'))
        r=export_split(self.preview(opts))
        for f in r.output_files:
            w=load_workbook(f)
            try:self.assertEqual(w.active['B2'].value,'001')
            finally:w.close()

    def test_excel_cached_error_result_cannot_be_hidden_by_blank_filter(self):
        o=SplitOptions(Path(__file__).parent/'fixtures/category_formula_cache.xlsx','错误结果',key_columns=(1,),output_dir=self.root)
        with self.assertRaisesRegex(ValueError,'B2'):export_split(self.preview(o))
