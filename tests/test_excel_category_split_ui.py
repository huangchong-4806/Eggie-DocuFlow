import os
os.environ.setdefault('QT_QPA_PLATFORM','offscreen')
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from PySide6.QtWidgets import QApplication
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from ui.excel_category_split import CategorySplitPage


class CategoryUiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app=QApplication.instance() or QApplication([])

    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.source=self.root/'source.xlsx'
        wb=Workbook();ws=wb.active;ws.title='明细'
        for row in [('供应商','金额'),('A',10),('B',20),('A',30)]:ws.append(row)
        ws.row_dimensions[3].hidden=True
        wb.save(self.source);wb.close()
        self.page=CategorySplitPage();self.page.show()
        self.addCleanup(self.page.close)

    def wait_idle(self):
        for _ in range(1000):
            self.app.processEvents();QTest.qWait(5)
            if not self.page.busy:return
        self.fail('后台任务未完成')

    def prepare(self, visible=False):
        self.page.source_edit.setText(str(self.source))
        QTest.mouseClick(self.page.load_button,Qt.LeftButton);self.wait_idle()
        self.page.output_edit.setText(str(self.root))
        QTest.mouseClick(self.page.next_button,Qt.LeftButton);self.wait_idle()
        self.assertEqual(self.page.steps.currentIndex(),1)
        self.page.column_list.item(0).setCheckState(Qt.Checked)
        with patch.object(self.page,'ask_visibility',return_value=visible):
            QTest.mouseClick(self.page.next_button,Qt.LeftButton);self.wait_idle()
        self.assertEqual(self.page.steps.currentIndex(),2)

    def test_cancel_confirmation_writes_nothing_and_retains_selection(self):
        self.prepare()
        self.page.category_table.item(1,0).setCheckState(Qt.Unchecked)
        with patch.object(self.page,'ask_confirmation',return_value=False):
            QTest.mouseClick(self.page.confirm_button,Qt.LeftButton)
        self.assertEqual(list(self.root.iterdir()),[self.source])
        self.assertEqual(self.page.preview.counts.exported,2)
        self.assertIn('未选',self.page.summary_label.text())

    def test_visible_choice_and_confirm_create_only_visible_records(self):
        self.prepare(visible=True)
        self.assertEqual(self.page.preview.counts.hidden,1)
        with patch.object(self.page,'ask_confirmation',return_value=True):
            QTest.mouseClick(self.page.confirm_button,Qt.LeftButton)
            self.wait_idle()
        self.assertIsNotNone(self.page.result)
        wb=load_workbook(self.page.result.output_files[0]);self.addCleanup(wb.close)
        self.assertEqual(list(wb.active.values),[('供应商','金额'),('A',10),('A',30)])
        self.assertFalse(self.page.confirm_button.isEnabled())

    def test_select_none_disables_execution(self):
        self.prepare()
        QTest.mouseClick(self.page.none_button,Qt.LeftButton)
        self.assertFalse(self.page.confirm_button.isEnabled())
        self.assertEqual(self.page.preview.counts.unselected,3)

    def test_back_change_invalidates_old_preview(self):
        self.prepare()
        QTest.mouseClick(self.page.back_button,Qt.LeftButton)
        self.page.trim_check.setChecked(True)
        self.assertIsNone(self.page.preview)
        self.assertFalse(self.page.confirm_button.isEnabled())

    def test_visibility_cancel_creates_no_preview_or_files(self):
        self.page.source_edit.setText(str(self.source));self.page.load_source();self.wait_idle()
        self.page.go_next();self.wait_idle();self.page.column_list.item(0).setCheckState(Qt.Checked)
        with patch.object(self.page,'ask_visibility',return_value=None):
            self.page.go_next();self.wait_idle()
        self.assertIsNone(self.page.preview)
        self.assertEqual(list(self.root.iterdir()),[self.source])

    def test_source_changed_after_column_selection_requires_reload(self):
        self.page.source_edit.setText(str(self.source));self.page.load_source();self.wait_idle()
        self.page.go_next();self.wait_idle();self.page.column_list.item(0).setCheckState(Qt.Checked)
        wb=load_workbook(self.source);wb.active['A1']='完全不同的字段';wb.save(self.source);wb.close()
        with patch.object(self.page,'ask_visibility',return_value=False):
            self.page.go_next();self.wait_idle()
        self.assertIsNone(self.page.preview)
        self.assertIn('变化',self.page.status_label.text())

    def test_double_confirm_click_cannot_start_duplicate_exports(self):
        self.prepare()
        with patch.object(self.page,'ask_confirmation',return_value=True):
            self.page.confirm_export();self.page.confirm_export();self.wait_idle()
        self.assertEqual(len([p for p in self.root.iterdir() if p.is_dir()]),1)
        self.assertEqual(self.page.result.exported_rows,3)

    def test_sheet_output_and_partial_selection_summary_matches_actual(self):
        self.page.source_edit.setText(str(self.source));self.page.load_source();self.wait_idle()
        self.page.go_next();self.wait_idle();self.page.column_list.item(0).setCheckState(Qt.Checked)
        self.page.mode_combo.setCurrentIndex(1)
        with patch.object(self.page,'ask_visibility',return_value=False):self.page.go_next();self.wait_idle()
        self.page.category_table.item(0,0).setCheckState(Qt.Unchecked)
        self.assertIn('导出 1 条',self.page.summary_label.text())
        self.assertIn('未选 2 条',self.page.summary_label.text())
        with patch.object(self.page,'ask_confirmation',return_value=True):self.page.confirm_export();self.wait_idle()
        w=load_workbook(self.page.result.output_files[0]);self.addCleanup(w.close)
        self.assertEqual(w.sheetnames,['B']);self.assertEqual(w.active['B2'].value,20)
